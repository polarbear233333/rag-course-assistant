import os
from io import BytesIO
from typing import Dict, List

import docx2txt
import fitz
import pytesseract
from PIL import Image
from PyPDF2 import PdfReader
from pptx import Presentation

try:
    from .config import DATA_DIR
except ImportError:
    from config import DATA_DIR


class DocumentLoader:
    def __init__(self, data_dir: str = DATA_DIR):
        self.data_dir = data_dir
        self.supported_formats = [".pdf", ".pptx", ".docx", ".txt"]

    def load_pdf(self, file_path: str) -> List[Dict]:
        pages = []
        try:
            reader = PdfReader(file_path)
            image_ocr_logged = False
            try:
                pdf_for_images = fitz.open(file_path)
            except Exception:
                pdf_for_images = None

            for page_num, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                ocr_texts = []

                if pdf_for_images is not None:
                    try:
                        fitz_page = pdf_for_images[page_num - 1]
                        for img in fitz_page.get_images():
                            xref = img[0]
                            pix = fitz_page.get_pixmap(xref=xref)
                            image = Image.open(BytesIO(pix.tobytes("png")))
                            ocr_result = pytesseract.image_to_string(image, lang="chi_sim+eng")
                            if ocr_result and ocr_result.strip():
                                ocr_texts.append(ocr_result.strip())
                    except Exception:
                        pass

                if ocr_texts:
                    if not image_ocr_logged:
                        print("已识别 PDF 图片中的文字")
                        image_ocr_logged = True
                    text = text + "\n【图片 OCR】\n" + "\n".join(ocr_texts)

                formatted_text = f"--- 第 {page_num} 页 ---\n{text}\n"
                pages.append({"text": formatted_text})

            if pdf_for_images is not None:
                pdf_for_images.close()
        except Exception as exc:
            print(f"加载 PDF 失败: {exc}")
        return pages

    def load_pptx(self, file_path: str) -> List[Dict]:
        slides = []
        try:
            presentation = Presentation(file_path)
            image_ocr_logged = False
            for slide_num, slide in enumerate(presentation.slides, 1):
                text_content = ""
                ocr_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_content += shape.text + "\n"
                    if hasattr(shape, "image"):
                        try:
                            image = Image.open(BytesIO(shape.image.blob))
                            ocr_result = pytesseract.image_to_string(image, lang="chi_sim+eng")
                            if ocr_result and ocr_result.strip():
                                ocr_texts.append(ocr_result.strip())
                        except Exception:
                            pass

                if ocr_texts:
                    if not image_ocr_logged:
                        print("已识别 PPT 图片中的文字")
                        image_ocr_logged = True
                    text_content = text_content + "\n【图片 OCR】\n" + "\n".join(ocr_texts)

                formatted_text = f"--- 幻灯片 {slide_num} ---\n{text_content}\n"
                slides.append({"text": formatted_text})
        except Exception as exc:
            print(f"加载 PPTX 失败: {exc}")
        return slides

    def load_docx(self, file_path: str) -> str:
        try:
            return docx2txt.process(file_path)
        except Exception as exc:
            print(f"加载 DOCX 失败: {exc}")
            return ""

    def load_txt(self, file_path: str) -> str:
        encodings = ["utf-8", "utf-8-sig", "gbk"]
        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as file:
                    return file.read()
            except UnicodeDecodeError:
                continue
            except Exception as exc:
                print(f"加载 TXT 失败: {exc}")
                return ""
        print(f"加载 TXT 失败: 无法识别编码 {file_path}")
        return ""

    def load_document(self, file_path: str) -> List[Dict[str, str]]:
        ext = os.path.splitext(file_path)[1].lower()
        filename = os.path.basename(file_path)
        documents = []

        if ext == ".pdf":
            pages = self.load_pdf(file_path)
            for page_idx, page_data in enumerate(pages, 1):
                documents.append(
                    {
                        "content": page_data["text"],
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": page_idx,
                    }
                )
        elif ext == ".pptx":
            slides = self.load_pptx(file_path)
            for slide_idx, slide_data in enumerate(slides, 1):
                documents.append(
                    {
                        "content": slide_data["text"],
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": slide_idx,
                    }
                )
        elif ext == ".docx":
            content = self.load_docx(file_path)
            if content:
                documents.append(
                    {
                        "content": content,
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": 0,
                    }
                )
        elif ext == ".txt":
            content = self.load_txt(file_path)
            if content:
                documents.append(
                    {
                        "content": content,
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": 0,
                    }
                )
        else:
            print(f"不支持的文件格式: {ext}")

        return documents

    def load_all_documents(self) -> List[Dict[str, str]]:
        if not os.path.exists(self.data_dir):
            print(f"数据目录不存在: {self.data_dir}")
            return []

        documents = []
        for root, _, files in os.walk(self.data_dir):
            for file in files:
                ext = os.path.splitext(file)[1].lower()
                if ext in self.supported_formats:
                    file_path = os.path.join(root, file)
                    print(f"正在加载: {file_path}")
                    doc_chunks = self.load_document(file_path)
                    if doc_chunks:
                        documents.extend(doc_chunks)
        return documents
